#!/usr/bin/env python3
"""Convert a PowerPoint (.pptx) deck into a .3df document.

  python pptx2_3df.py deck.pptx [deck.3df] [--dpi 200] [--soffice PATH]

Google Slides: File > Download > Microsoft PowerPoint (.pptx), then run this.

How it works
  1. LibreOffice renders every visible slide; each becomes a full-page image.
     This keeps fonts, shapes, charts, tables and backgrounds exactly as drawn.
  2. Animated GIFs and PowerPoint 3D models are found in the slide XML and laid
     over the page at the same position, so they play and rotate in 3DF.
     The rendered page region underneath becomes each element's fallback.

Requires: LibreOffice, and `pip install python-pptx lxml pymupdf pillow`.
"""
import argparse
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation

import threedf

EMU_PER_MM = 36000
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
R_EMBED = f"{{{NS['r']}}}embed"


def find_soffice(explicit=None):
    candidates = [explicit, shutil.which("soffice"), shutil.which("libreoffice"),
                  "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                  r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]
    for c in candidates:
        if c and Path(c).exists():
            return c
    sys.exit("LibreOffice not found. Install it from libreoffice.org, or pass --soffice PATH.")


def render_pdf(pptx, workdir, soffice):
    profile = Path(workdir, "lo_profile").resolve().as_uri()
    subprocess.run([soffice, f"-env:UserInstallation={profile}", "--headless",
                    "--convert-to", "pdf", "--outdir", str(workdir), str(pptx)],
                   check=True, capture_output=True, timeout=600)
    pdf = Path(workdir, Path(pptx).stem + ".pdf")
    if not pdf.exists():
        sys.exit("LibreOffice ran but produced no PDF. Check the deck opens in LibreOffice Impress.")
    return pdf


# ---------- geometry ----------
def _xfrm(el):
    """Return (x, y, cx, cy, rot, flipped) for a shape's own a:xfrm or p:xfrm."""
    x = el.find("p:spPr/a:xfrm", NS)
    if x is None:
        x = el.find("p:xfrm", NS)
    if x is None:
        x = el.find("p:grpSpPr/a:xfrm", NS)
    if x is None:
        return None
    off, ext = x.find("a:off", NS), x.find("a:ext", NS)
    rot = int(x.get("rot", "0"))
    flipped = x.get("flipH") == "1" or x.get("flipV") == "1"
    return [int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy")), rot, flipped]


def to_slide_space(el, box):
    """Map a box from its group's child space up through every enclosing p:grpSp."""
    x, y, cx, cy = box
    rotated = False
    node = el.getparent()
    while node is not None:
        if node.tag == f"{{{NS['p']}}}grpSp":
            gx = node.find("p:grpSpPr/a:xfrm", NS)
            off, ext = gx.find("a:off", NS), gx.find("a:ext", NS)
            if gx is None or gx.find("a:off", NS) is None or gx.find("a:ext", NS) is None:
                node = node.getparent()
                continue
            choff = gx.find("a:chOff", NS)
            chext = gx.find("a:chExt", NS)
            if choff is None:
                choff = gx.find("a:off", NS)
            if chext is None:
                chext = gx.find("a:ext", NS)
            sx = int(ext.get("cx")) / max(1, int(chext.get("cx")))
            sy = int(ext.get("cy")) / max(1, int(chext.get("cy")))
            x = int(off.get("x")) + (x - int(choff.get("x"))) * sx
            y = int(off.get("y")) + (y - int(choff.get("y"))) * sy
            cx, cy = cx * sx, cy * sy
            rotated |= gx.get("rot", "0") != "0" or gx.get("flipH") == "1" or gx.get("flipV") == "1"
        node = node.getparent()
    return (x, y, cx, cy), rotated


def in_fallback(el):
    node = el.getparent()
    while node is not None:
        if node.tag == f"{{{NS['mc']}}}Fallback":
            return True
        node = node.getparent()
    return False


# ---------- slide scanning ----------
def scan_slide(slide):
    """Yield dicts for animated GIFs and 3D models on one slide, plus warnings."""
    found, warnings = [], []
    root = slide.element

    for pic in root.iter(f"{{{NS['p']}}}pic"):
        if in_fallback(pic):
            continue
        blip = pic.find("p:blipFill/a:blip", NS)
        if blip is None or blip.get(R_EMBED) is None:
            continue
        part = slide.part.related_part(blip.get(R_EMBED))
        blob = part.blob
        if not blob[:6] in (b"GIF87a", b"GIF89a"):
            continue
        try:
            if getattr(Image.open(io.BytesIO(blob)), "n_frames", 1) < 2:
                continue
        except Exception:
            continue
        name = pic.find("p:nvPicPr/p:cNvPr", NS)
        label = (name.get("descr") or name.get("name") or "Animated image") if name is not None else "Animated image"
        xf = _xfrm(pic)
        if xf is None:
            continue
        src_rect = pic.find("p:blipFill/a:srcRect", NS)
        cropped = src_rect is not None and any(src_rect.get(k, "0") != "0" for k in "ltrb")
        box, grp_rot = to_slide_space(pic, xf[:4])
        if xf[4] or xf[5] or grp_rot or cropped:
            warnings.append(f'"{label}" is rotated, flipped or cropped, so it stays a still image')
            continue
        found.append({"kind": "gif", "blob": blob, "box": box, "alt": label})

    for m in root.iter():
        if not isinstance(m.tag, str) or not m.tag.endswith("}model3d"):
            continue
        frame = m
        while frame is not None and frame.tag != f"{{{NS['p']}}}graphicFrame":
            frame = frame.getparent()
        rid = m.get(R_EMBED)
        if frame is None or rid is None:
            continue
        blob = slide.part.related_part(rid).blob
        if blob[:4] != b"glTF":
            warnings.append("A 3D model isn't glTF binary, so it stays a still image")
            continue
        xf = _xfrm(frame)
        if xf is None:
            continue
        box, grp_rot = to_slide_space(frame, xf[:4])
        cnv = frame.find("p:nvGraphicFramePr/p:cNvPr", NS)
        label = (cnv.get("descr") or cnv.get("name") or "3D model") if cnv is not None else "3D model"
        if grp_rot or xf[4] or xf[5]:
            warnings.append(f'"{label}" is rotated or flipped; placed without rotation')
        found.append({"kind": "model", "blob": blob, "box": box, "alt": label})
    return found, warnings


def edge_color(img):
    """Most common colour around a crop's border, used as the 3D viewport background."""
    w, h = img.size
    px = img.convert("RGB").load()
    border = [px[x, y] for x in range(0, w, max(1, w // 40)) for y in (0, h - 1)]
    border += [px[x, y] for y in range(0, h, max(1, h // 40)) for x in (0, w - 1)]
    r, g, b = Counter(border).most_common(1)[0][0]
    return f"#{r:02x}{g:02x}{b:02x}"


def convert(pptx, out, dpi=200, soffice=None):
    if not Path(pptx).is_file():
        sys.exit(f"{pptx} not found")
    if Path(pptx).suffix.lower() != ".pptx":
        sys.exit("Only .pptx is supported. Save older .ppt files as .pptx first.")
    prs = Presentation(pptx)
    W_mm = (prs.slide_width or 12192000) / EMU_PER_MM    # default 16:9 if the deck omits a size
    H_mm = (prs.slide_height or 6858000) / EMU_PER_MM
    visible = [s for s in prs.slides if s.element.get("show") != "0"]

    with tempfile.TemporaryDirectory() as tmp:
        pdf = render_pdf(pptx, tmp, find_soffice(soffice))
        pdoc = pymupdf.open(pdf)
        if len(pdoc) != len(visible):
            sys.exit(f"LibreOffice rendered {len(pdoc)} pages but the deck has {len(visible)} visible slides.")

        root = Path(tmp, "doc")
        (root / "assets").mkdir(parents=True)
        pages, all_warnings = [], []
        for i, (slide, page) in enumerate(zip(visible, pdoc), 1):
            pix = page.get_pixmap(dpi=dpi)
            bg = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            bg_path = f"assets/slide{i}.png"
            bg.save(root / bg_path, optimize=True)
            px_per_mm = bg.width / W_mm
            title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
            elements = [{"type": "image", "x": 0, "y": 0, "w": W_mm, "h": H_mm, "src": bg_path,
                         "alt": title or f"Slide {i}", "z": 0}]

            found, warnings = scan_slide(slide)
            all_warnings += [f"Slide {i}: {w}" for w in warnings]
            for k, f in enumerate(found, 1):
                x, y, cx, cy = (v / EMU_PER_MM for v in f["box"])
                l, t = max(0, round(x * px_per_mm)), max(0, round(y * px_per_mm))
                r, b = min(bg.width, round((x + cx) * px_per_mm)), min(bg.height, round((y + cy) * px_per_mm))
                if cx <= 0 or cy <= 0 or r - l < 2 or b - t < 2:
                    all_warnings.append(f'Slide {i}: "{f["alt"]}" is off the slide or has no size, skipped')
                    continue
                crop = bg.crop((l, t, r, b))
                fb = f"assets/slide{i}_{k}_fallback.png"
                crop.save(root / fb)
                el = {"x": round(x, 2), "y": round(y, 2), "w": round(cx, 2), "h": round(cy, 2),
                      "alt": f["alt"], "fallback": fb, "z": k}
                if f["kind"] == "gif":
                    src = f"assets/slide{i}_{k}.gif"
                    el.update(type="image", src=src)
                else:
                    src = f"assets/slide{i}_{k}.glb"
                    el.update(type="model", src=src, background=edge_color(crop))
                (root / src).write_bytes(f["blob"])
                elements.append(el)
            pages.append({"elements": elements})

        doc = {"3df": threedf.SPEC_VERSION,
               "title": prs.core_properties.title or Path(pptx).stem,
               "author": prs.core_properties.author or "",
               "page": {"width_mm": round(W_mm, 2), "height_mm": round(H_mm, 2)},
               "pages": pages}
        (root / "document.json").write_text(json.dumps(doc, indent=2))
        threedf.pack(root, out)
    return len(pages), all_warnings


def main():
    ap = argparse.ArgumentParser(description="Convert PowerPoint (.pptx) to .3df")
    ap.add_argument("pptx")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--dpi", type=int, default=200, help="slide render resolution (default 200)")
    ap.add_argument("--soffice", help="path to the LibreOffice soffice executable")
    a = ap.parse_args()
    out = a.out or str(Path(a.pptx).with_suffix(".3df"))
    n, warnings = convert(a.pptx, out, a.dpi, a.soffice)
    for w in warnings:
        print("note:", w)
    print(f"wrote {out} ({n} pages)")


if __name__ == "__main__":
    main()
