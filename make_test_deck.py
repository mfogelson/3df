"""Builds test_deck.pptx: GIF, 3D model, hidden slide, grouped GIF, rotated GIF."""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI

A="sample/assets/"
prs=Presentation(); prs.slide_width=Mm(338.67); prs.slide_height=Mm(190.5)
blank=prs.slide_layouts[5]  # title only
def slide(t):
    s=prs.slides.add_slide(blank); s.shapes.title.text=t; return s

s=slide("Panel deployment"); s.shapes.add_picture(A+"deploy.gif",Mm(60),Mm(55),Mm(220),Mm(110))

s=slide("Truss model")
pic=s.shapes.add_picture(A+"truss_fallback.png",Mm(40),Mm(50),Mm(260),Mm(130))
glb=Part(PackURI("/ppt/media/model3d1.glb"),"model/gltf-binary",prs.part.package,open(A+"truss.glb","rb").read())
rid=s.part.relate_to(glb,"http://schemas.microsoft.com/office/2017/06/relationships/model3d")
P="http://schemas.openxmlformats.org/presentationml/2006/main"
xml=f'''<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p="{P}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<mc:Choice xmlns:am3d="http://schemas.microsoft.com/office/drawing/2017/model3d" Requires="am3d">
<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="90" name="3D Model 1" descr="Four-bay truss"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
<p:xfrm><a:off x="{Mm(40)}" y="{Mm(50)}"/><a:ext cx="{Mm(260)}" cy="{Mm(130)}"/></p:xfrm>
<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/drawing/2017/model3d"><am3d:model3d r:embed="{rid}"/></a:graphicData></a:graphic></p:graphicFrame>
</mc:Choice><mc:Fallback/></mc:AlternateContent>'''
ac=etree.fromstring(xml); pel=pic._element; pel.addprevious(ac); ac[-1].append(pel)

s=slide("Hidden slide"); s._element.set("show","0")

s=slide("Grouped GIF"); g=s.shapes.add_group_shape(); g.shapes.add_picture(A+"deploy.gif",Mm(30),Mm(60),Mm(120),Mm(60)); g.shapes.add_picture(A+"deploy_fallback.png",Mm(180),Mm(60),Mm(120),Mm(60))

s=slide("Rotated GIF"); p=s.shapes.add_picture(A+"deploy.gif",Mm(80),Mm(60),Mm(160),Mm(80)); p.rotation=10
prs.save("test_deck.pptx"); print("ok")
